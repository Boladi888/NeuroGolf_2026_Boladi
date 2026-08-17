import csv
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl

from tools.neurogolf_sync_task_state import infer_changed_task_entries, sync_task_state


def _write_zip(path: Path, entries: dict[str, bytes]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with ZipFile(path, "w", ZIP_DEFLATED) as zf:
        for name, data in entries.items():
            zf.writestr(name, data)


def _seed_task_state(repo: Path, rows: list[dict[str, str]] | None = None) -> tuple[Path, Path]:
    task_dir = repo / "The_400_Tasks"
    task_dir.mkdir(parents=True)
    headers = [
        "task",
        "task_id",
        "status",
        "best_known_status",
        "last_attempted_at",
        "current_cost",
        "current_src",
        "current_points",
        "macs",
        "memory",
        "params",
        "verdict",
        "next_action",
        "local_payload_path",
        "reference",
        "notes",
        "structural_attempt",
        "unlock_potential",
        "structural_kaggle_delta",
    ]
    row = {
        "task": "task013",
        "task_id": "13",
        "status": "WIP",
        "best_known_status": "old",
        "last_attempted_at": "",
        "current_cost": "58000",
        "current_src": "old_src",
        "current_points": "14",
        "macs": "0",
        "memory": "57900",
        "params": "100",
        "verdict": "old verdict",
        "next_action": "manual action should stay",
        "local_payload_path": "old.onnx",
        "reference": "old ref",
        "notes": "old note",
        "structural_attempt": "not_yet",
        "unlock_potential": "HIGH",
        "structural_kaggle_delta": "",
    }
    rows = rows or [row]
    csv_path = task_dir / "task_state.csv"
    with csv_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)

    xlsx_path = task_dir / "NeuroGolf_Task_Status.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tasks"
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header, "") for header in headers])
    wb.save(xlsx_path)
    return csv_path, xlsx_path


def test_infer_changed_task_entries_uses_zip_diff_when_metadata_is_stacky(tmp_path: Path):
    old_zip = tmp_path / "old.zip"
    new_zip = tmp_path / "new.zip"
    _write_zip(old_zip, {"task001.onnx": b"old", "task002.onnx": b"same"})
    _write_zip(new_zip, {"task001.onnx": b"new", "task002.onnx": b"same"})

    changed = infer_changed_task_entries(old_zip, new_zip, ["megastack153-153-task-stack.onnx"])

    assert changed == ["task001.onnx"]


def test_sync_task_state_patches_matrix_and_cost_cache_without_enrich(tmp_path: Path):
    csv_path, xlsx_path = _seed_task_state(tmp_path)
    package_path = tmp_path / "submissions" / "submission.zip"
    package_path.parent.mkdir()
    package_path.write_bytes(b"not used because scores are injected")

    result = sync_task_state(
        repo_root=tmp_path,
        package_path=package_path,
        changed_entries=["task013.onnx"],
        candidate_id="task013_rewrite_modgather_t_006ff1dc_on_6448_40",
        public_score=6451.44,
        previous_score=6448.40,
        delta=3.04,
        completed_at="2026-06-01T23:42:30Z",
        mark_shipped=True,
        run_enrich=False,
        precomputed_scores={
            13: {
                "entry": "task013.onnx",
                "score_status": "ok",
                "macs": 0,
                "memory": 2708,
                "params": 66,
                "current_cost": 2774,
                "current_points": 17.071954399125225,
            }
        },
    )

    assert result["updated_task_ids"] == [13]
    assert (tmp_path / "The_400_Tasks" / "latest_package_costs_2026-06-01.csv").exists()

    with csv_path.open(encoding="utf-8", newline="") as handle:
        rows = list(csv.DictReader(handle))
    row = next(item for item in rows if item["task"] == "task013")
    assert row["status"] == "shipped"
    assert row["best_known_status"] == "kaggle_confirmed_2026-06-01"
    assert row["current_cost"] == "2774"
    assert row["current_src"] == "task013_rewrite_modgather_t_006ff1dc"
    assert row["memory"] == "2708"
    assert row["params"] == "66"
    assert "Kaggle COMPLETE: 6448.4 -> 6451.44" in row["verdict"]
    assert row["next_action"] == "manual action should stay"
    assert "old note | Shipped 2026-06-01" in row["notes"]

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["Tasks"]
    headers = [cell.value for cell in ws[1]]
    task_col = headers.index("task") + 1
    task_row = next(row_idx for row_idx in range(2, ws.max_row + 1) if ws.cell(row_idx, task_col).value == "task013")
    values = {header: ws.cell(task_row, idx + 1).value for idx, header in enumerate(headers)}
    assert str(values["current_cost"]) == "2774"
    assert values["status"] == "shipped"


def test_multi_task_package_sync_preserves_non_primary_sources(tmp_path: Path):
    base = {
        "status": "shipped",
        "best_known_status": "kaggle_confirmed_2026-06-01",
        "last_attempted_at": "",
        "current_points": "15",
        "macs": "0",
        "verdict": "old verdict",
        "next_action": "manual action should stay",
        "local_payload_path": "old.onnx",
        "reference": "old ref",
        "structural_attempt": "shipped",
        "unlock_potential": "LOW",
    }
    csv_path, _ = _seed_task_state(
        tmp_path,
        [
            {
                **base,
                "task": "task138",
                "task_id": "138",
                "current_cost": "14838",
                "current_src": "task138_coloridx_cropcollapse_v7",
                "memory": "14596",
                "params": "242",
                "notes": "old task138 note",
                "structural_kaggle_delta": "+1.38",
            },
            {
                **base,
                "task": "task213",
                "task_id": "213",
                "current_cost": "55705",
                "current_src": "old_task213",
                "memory": "55000",
                "params": "705",
                "notes": "old task213 note",
                "structural_kaggle_delta": "",
            },
        ],
    )
    package_path = tmp_path / "submissions" / "submission.zip"
    package_path.parent.mkdir()
    package_path.write_bytes(b"not used because scores are injected")

    sync_task_state(
        repo_root=tmp_path,
        package_path=package_path,
        changed_entries=["task138.onnx", "task213.onnx"],
        candidate_id="task213_v8_t_2193_on_6451_71",
        primary_task_id=213,
        public_score=6454.94,
        previous_score=6451.71,
        delta=3.23,
        completed_at="2026-06-02T01:11:00Z",
        mark_shipped=True,
        run_enrich=False,
        precomputed_scores={
            138: {
                "entry": "task138.onnx",
                "score_status": "ok",
                "macs": 0,
                "memory": 12898,
                "params": 306,
                "current_cost": 13204,
                "current_points": 15.5,
            },
            213: {
                "entry": "task213.onnx",
                "score_status": "ok",
                "macs": 0,
                "memory": 2122,
                "params": 71,
                "current_cost": 2193,
                "current_points": 17.3,
            },
        },
    )

    with csv_path.open(encoding="utf-8", newline="") as handle:
        rows = {row["task"]: row for row in csv.DictReader(handle)}
    assert rows["task213"]["current_src"] == "task213_v8_t_2193"
    assert rows["task213"]["structural_kaggle_delta"] == "+3.23"
    assert "Shipped 2026-06-02: task213_v8_t_2193" in rows["task213"]["notes"]
    assert rows["task138"]["current_src"] == "task138_coloridx_cropcollapse_v7"
    assert rows["task138"]["structural_kaggle_delta"] == "+1.38"
    assert "Package sync 2026-06-02" in rows["task138"]["notes"]
    assert "Shipped 2026-06-02: task213_v8_t_2193" not in rows["task138"]["notes"]
