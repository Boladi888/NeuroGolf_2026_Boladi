"""Merge The_400_Tasks/task_state.csv into NeuroGolf_Task_Status.xlsx.

Outputs:
  - The_400_Tasks/NeuroGolf_Task_Status.xlsx (rewritten with merged Tasks sheet,
    color coding, frozen header row).
  - The_400_Tasks/task_state.csv (regenerated as a flat CSV mirror of the
    Tasks sheet for diffability / programmatic access).

Adds new columns:
  - last_attempted_at: ISO 8601 datetime of last *active* attempt (not just inspected).
    Free text. Populate manually or via tooling that updates this field at the end
    of each focused attempt session.
  - time_spent_attempting: cumulative duration spent attempting this task. Free
    text (e.g., "2h30m", "45m") so humans/agents can update by hand.

Idempotent: re-running with an updated xlsx + CSV merges fresh values back in
without losing manual edits. Cost/points/src/split columns: the CSV is the live source of truth (captain reconciles it per ship, 2026-06-07); workflow/classification columns: the xlsx wins.

Usage:
    python tools/merge_task_status.py [--xlsx PATH] [--csv PATH]
"""
from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import CellIsRule, FormulaRule, ColorScaleRule
from openpyxl.utils import get_column_letter


REPO = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = REPO / "The_400_Tasks" / "NeuroGolf_Task_Status.xlsx"
DEFAULT_CSV = REPO / "The_400_Tasks" / "task_state.csv"


# Final column order in the merged Tasks sheet. Group by purpose.
# (header, source ['xlsx'|'csv'|'new'], description)
COLUMNS = [
    # --- Identity ---
    ("task", "xlsx", "Canonical task ID string (taskNNN, zero-padded)"),
    ("task_id", "xlsx", "Integer 0-400"),
    # --- Classification ---
    ("family", "xlsx", "DeepSeek-classified semantic family (sparse — fill in)"),
    ("transformations", "xlsx", "Pipe-separated primitive list"),
    ("complexity_1_10", "xlsx", "DeepSeek complexity estimate"),
    ("grid_resized", "xlsx", "Output shape != input shape"),
    # --- Workflow (most important for humans) ---
    ("status", "xlsx", "Workflow state: unassigned | WIP | blocked | recipe-ready | locked-in | submitted | done | UNSOLVED | big-delta | review-ext"),
    ("best_known_status", "csv", "Per-task scoring verdict (from task_state.csv)"),
    ("priority", "xlsx", "H/M/L or numeric"),
    ("assigned_to", "xlsx", "Agent / branch on this task"),
    # --- New tracking columns ---
    ("last_attempted_at", "new", "ISO 8601 datetime of last *active attempt* (not just looking)"),
    ("time_spent_attempting", "new", "Cumulative duration on attempts (e.g., '2h30m')"),
    # --- Cost / scoring ---
    ("current_cost", "csv", "Our best per-task ONNX cost (lower = better) — CSV is the live reconcile target (flipped 2026-06-07; captain reconciles CSV per ship)"),
    ("current_src", "csv", "Source kernel/strategy that won this task in our blend (flipped 2026-06-07)"),
    ("current_points", "csv", "Per-task contribution to LB score (flipped 2026-06-07)"),
    ("best_ext_cost", "xlsx", "Cheapest external candidate ever observed"),
    ("best_ext_src", "xlsx", "Competitor producing that minimum"),
    ("ext_savings_pct", "xlsx", "(current-best_ext)/current x 100 — >10% = worth investigating"),
    ("floor_proxy", "xlsx", "Likely cheapest legal op chain (sparse)"),
    ("gap_to_floor", "xlsx", "current_cost - floor_proxy (sparse)"),
    ("macs", "xlsx", "ONNX op-profile MACs"),
    ("memory", "csv", "ONNX op-profile memory (flipped 2026-06-07)"),
    ("params", "csv", "ONNX op-profile params (flipped 2026-06-07)"),
    # --- Verdict / actions (free text) ---
    ("verdict", "csv", "Human-readable verdict line (from CSV)"),
    ("next_action", "csv", "What should happen next (from CSV)"),
    # --- Recipe pointers ---
    ("recipe_exists", "xlsx", "Bool — Research_Archive/notes/recipe-taskNNN-*.md present"),
    ("recipe_hint", "xlsx", "Short recipe summary if present"),
    ("recipe_path", "xlsx", "Path to recipe note"),
    # --- Artifact links ---
    ("candidate_zips", "csv", "Comma-separated candidate-zip names containing this task"),
    ("local_payload_path", "csv", "Path to local ONNX selected for this task"),
    ("reference", "csv", "Pointer to session log / digest / recipe"),
    # --- Batch metadata ---
    ("n_batch_runs", "xlsx", "Batch experiments that touched this task"),
    ("last_batch", "xlsx", "Newest batch JSON filename"),
    ("last_batch_reviewed", "xlsx", "Whether .reviewed sibling exists"),
    ("best_train_acc", "xlsx", "Best train acc from external candidates"),
    ("best_test_acc", "xlsx", "Best test acc from external candidates"),
    ("n_external_candidates", "xlsx", "Distinct competitor sources for this task"),
    # --- Free text ---
    ("notes", "merged", "Free text — gotchas, history, watchlist. Merged from xlsx.notes + csv.notes."),
    # --- Auto-enriched relational columns (regenerated by tools/enrich_task_state.py) ---
    ("ops_count", "xlsx", "Total nodes in current ONNX (sortable int)"),
    ("op_signature", "xlsx", "Top 3 ops by count, format 'Op:N|Op:N|Op:N' (groupable for batch fixes)"),
    ("unique_ops", "xlsx", "Sorted '|'-joined unique op_types in current ONNX"),
    ("graph_name", "xlsx", "ONNX graph.name (often hints at builder)"),
    ("producer", "xlsx", "ONNX producer_name (often hints at origin)"),
    ("similar_tasks", "xlsx", "Comma-joined taskNNN list of tasks sharing same family+transformations"),
    ("similar_task_count", "xlsx", "Count of similar_tasks (sortable int)"),
    ("current_src_shared_with", "xlsx", "Count of OTHER tasks using same current_src (high = cascade opportunity)"),
    ("solved_status", "xlsx", "Enum: Solved | NotSolved | HiddenBad"),
    ("hidden_bad_history", "xlsx", "Yes/No — task confirmed in silent-fail-trap"),
    ("live_arch", "xlsx", "Enum: lookup | lookup_confirmed_zero | trained_nn | structural. Auto-detected from current ONNX initializer/op signatures."),
    ("cherry13b_tested", "manual", "Manual enum; default not_yet. Suggested values: not_yet | no_delta | regresses | passed | n/a."),
    ("structural_attempt", "manual", "Manual enum/text; default not_yet. Updated when a structural-rule worker is assigned or shipped."),
    ("unlock_potential", "xlsx", "Computed: HIGH if live_arch=lookup and cherry13b_tested is not_yet/regresses; otherwise LOW."),
    ("structural_kaggle_delta", "manual", "Manual Kaggle delta from structural-rule ship/probe, filled after submission."),
    ("prototype_path", "xlsx", "Semicolon-joined task-local prototype artifact paths (.py/.onnx) if any exist."),
    ("team_split", "csv", "OnnxlyFans split: boladi | aditya | deferred (added 2026-06-07 — without this entry a merge run DROPS the column)"),
]


# Color palette
FILL_HEADER = PatternFill("solid", fgColor="1F4E78")  # dark blue
FONT_HEADER = Font(color="FFFFFF", bold=True)

# Family palette (sparse — used when family is populated)
FAMILY_COLORS = {
    "Pattern_Recognition": "FFE4B5",      # moccasin
    "Spatial_and_Geometric": "B5D9FF",    # light blue
    "Object_Based": "B5FFD9",             # light green
    "Color_and_Logical": "FFB5D9",        # light pink
    "Counting_and_Comparison": "E0B5FF",  # light purple
    "Movement_and_Gravity": "FFD9B5",     # light orange
    "Symmetry_and_Reflection": "B5FFE0",  # light cyan
}

# Status palette
STATUS_COLORS = {
    "done": "C6EFCE",          # green
    "locked-in": "C6EFCE",
    "submitted": "B5D9FF",     # blue
    "recipe-ready": "FFEB9C",  # yellow
    "WIP": "FFD966",           # gold
    "blocked": "FFC7CE",       # red
    "UNSOLVED": "F4B084",      # orange
    "big-delta": "FFEB9C",
    "review-ext": "D9E1F2",    # light blue
}

# solved_status palette
SOLVED_STATUS_COLORS = {
    "Solved": "C6EFCE",      # green
    "NotSolved": "F2F2F2",   # light gray
    "HiddenBad": "FFC7CE",   # red — danger
}


# hidden_bad_history palette
HIDDEN_BAD_COLORS = {
    "Yes": "FFC7CE",   # red
    "No": "F2F2F2",    # gray
}

LIVE_ARCH_COLORS = {
    "lookup": "FFC7CE",                 # red — action needed
    "lookup_confirmed_zero": "D9EAD3",  # green — already ruled out
    "trained_nn": "F4CCCC",             # red/pink — caution
    "structural": "F2F2F2",             # gray
}

CHERRY13B_COLORS = {
    "not_yet": "FFEB9C",
    "regresses": "C6EFCE",
    "regressed": "C6EFCE",
    "no_delta": "F2F2F2",
    "passed": "C6EFCE",
    "n/a": "D9EAD3",
}

UNLOCK_POTENTIAL_COLORS = {
    "HIGH": "FFC7CE",
    "LOW": "F2F2F2",
}


# Best_known_status palette
BEST_KNOWN_COLORS = {
    "organic_strict_clean": "C6EFCE",
    "tested_public_positive_current_best": "C6EFCE",
    "tested_public_positive_current_best_stack_member": "DDEBF7",
    "tested_public_positive_base_lineage": "DDEBF7",
    "tested_public_positive_isolated": "DDEBF7",
    "tested_non_improving": "FFEB9C",
    "queued_candidate_member": "F4E1B6",
    "no_current_replacement_record": "F2F2F2",
    "avoid_or_retest": "FFC7CE",
    "documentation_only_solved": "E2D5F0",
}


def load_csv(path: Path) -> dict[int, dict]:
    """Load task_state.csv into {task_id: row_dict}.

    Accepts both formats:
      - new: `task='task001', task_id=1, ...`
      - legacy: `task='001', ...` (old task_state.csv format pre-merge)
    """
    out: dict[int, dict] = {}
    if not path.exists():
        return out
    with path.open(encoding="utf-8", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            tid = None
            # Prefer task_id column
            tid_str = row.get("task_id")
            if tid_str:
                try:
                    tid = int(tid_str)
                except ValueError:
                    tid = None
            if tid is None:
                # Fall back to `task` (int or 'taskNNN')
                raw = row.get("task", "")
                if raw.startswith("task"):
                    raw = raw[4:]
                try:
                    tid = int(raw)
                except (ValueError, TypeError):
                    continue
            out[tid] = row
    return out


def load_xlsx(path: Path) -> tuple[dict[int, dict], dict]:
    """Load the existing Tasks sheet into {task_id: row_dict}.
    Also return a dict of (sheet_name -> rows) for non-Tasks sheets so we can
    preserve Summary, Floor_Proxy, Legend across rewrites.
    """
    if not path.exists():
        return {}, {}
    wb = openpyxl.load_workbook(path, data_only=True)
    rows: dict[int, dict] = {}
    other_sheets: dict[str, list[tuple]] = {}
    for sn in wb.sheetnames:
        ws = wb[sn]
        if sn == "Tasks":
            headers = [c.value for c in ws[1]]
            for row in ws.iter_rows(min_row=2, values_only=True):
                d = dict(zip(headers, row))
                tid = d.get("task_id")
                if isinstance(tid, int):
                    rows[tid] = d
        elif not sn.startswith("Schema"):
            other_sheets[sn] = [tuple(r) for r in ws.iter_rows(values_only=True)]
    return rows, other_sheets


def merge_rows(xlsx_rows: dict[int, dict], csv_rows: dict[int, dict]) -> list[dict]:
    """Merge per-task rows from both sources by task_id.

    For each of the 401 tasks (0..400):
      - Start with xlsx row (if any).
      - Overlay CSV columns (best_known_status, verdict, next_action, etc.)
        — CSV takes precedence for csv-source columns.
      - Preserve last_attempted_at, time_spent_attempting from existing xlsx
        (since CSV doesn't have them).
      - Merge notes (concatenate non-empty).
    """
    merged: list[dict] = []
    all_ids = sorted(set(xlsx_rows.keys()) | set(csv_rows.keys()) | set(range(401)))
    for tid in all_ids:
        if tid > 400:
            continue
        x = xlsx_rows.get(tid, {}) or {}
        c = csv_rows.get(tid, {}) or {}
        row = {}
        for col, src, _desc in COLUMNS:
            val = None
            if src == "xlsx":
                val = x.get(col)
                if (val is None or val == "") and c.get(col):
                    val = c.get(col)
            elif src == "csv":
                val = c.get(col)
                if (val is None or val == "") and x.get(col):
                    val = x.get(col)
            elif src == "new":
                # Preserve any value already in xlsx (manual edit), else blank.
                val = x.get(col) or c.get(col)
            elif src == "manual":
                # Manual fields can be edited from either surface; prefer xlsx
                # when populated because it preserves Excel-side edits.
                val = x.get(col) or c.get(col)
            elif src == "merged":
                # notes: combine xlsx + csv if both populated
                xn = x.get(col) or ""
                cn = c.get(col) or ""
                if xn and cn and xn != cn:
                    val = f"{xn} | {cn}"
                else:
                    val = xn or cn or None
            row[col] = val
        row["task"] = f"task{tid:03d}"
        row["task_id"] = tid
        merged.append(row)
    return merged


def write_xlsx(path: Path, merged: list[dict], other_sheets: dict) -> None:
    wb = Workbook()
    # Default sheet — replace with Tasks
    ws = wb.active
    ws.title = "Tasks"
    headers = [c[0] for c in COLUMNS]
    for ci, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.fill = FILL_HEADER
        cell.font = FONT_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for ri, row in enumerate(merged, start=2):
        for ci, h in enumerate(headers, start=1):
            ws.cell(row=ri, column=ci, value=row.get(h))
    # Freeze header
    ws.freeze_panes = "C2"
    # Column widths
    widths = {
        "task": 10, "task_id": 8, "family": 22, "transformations": 36,
        "complexity_1_10": 8, "grid_resized": 8, "status": 14,
        "best_known_status": 32, "priority": 8, "assigned_to": 16,
        "last_attempted_at": 22, "time_spent_attempting": 16,
        "current_cost": 13, "current_src": 18, "current_points": 11,
        "best_ext_cost": 13, "best_ext_src": 18, "ext_savings_pct": 11,
        "floor_proxy": 12, "gap_to_floor": 12,
        "macs": 11, "memory": 11, "params": 9,
        "verdict": 48, "next_action": 44,
        "recipe_exists": 8, "recipe_hint": 22, "recipe_path": 32,
        "candidate_zips": 30, "local_payload_path": 28, "reference": 32,
        "n_batch_runs": 8, "last_batch": 26, "last_batch_reviewed": 10,
        "best_train_acc": 10, "best_test_acc": 10, "n_external_candidates": 10,
        "notes": 36,
        # relational columns
        "ops_count": 9, "op_signature": 34, "unique_ops": 40,
        "graph_name": 26, "producer": 16,
        "similar_tasks": 50, "similar_task_count": 8,
        "current_src_shared_with": 10,
        "solved_status": 12, "hidden_bad_history": 8,
        "live_arch": 20, "cherry13b_tested": 16,
        "structural_attempt": 18, "unlock_potential": 16,
        "structural_kaggle_delta": 18,
        "prototype_path": 44,
    }
    for ci, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = widths.get(h, 14)

    last_row = 1 + len(merged)

    # Color coding
    def col_idx(name: str) -> int:
        return headers.index(name) + 1

    def col_letter(name: str) -> str:
        return get_column_letter(col_idx(name))

    # 1. Family colors — fill cells in the family column by category
    fam_idx = col_idx("family")
    for ri, row in enumerate(merged, start=2):
        fam = row.get("family")
        if fam and fam in FAMILY_COLORS:
            ws.cell(row=ri, column=fam_idx).fill = PatternFill("solid", fgColor=FAMILY_COLORS[fam])

    # 2. Status colors
    status_idx = col_idx("status")
    for ri, row in enumerate(merged, start=2):
        s = row.get("status")
        if s and s in STATUS_COLORS:
            ws.cell(row=ri, column=status_idx).fill = PatternFill("solid", fgColor=STATUS_COLORS[s])

    # 3. best_known_status colors
    bks_idx = col_idx("best_known_status")
    for ri, row in enumerate(merged, start=2):
        s = row.get("best_known_status")
        if s and s in BEST_KNOWN_COLORS:
            ws.cell(row=ri, column=bks_idx).fill = PatternFill("solid", fgColor=BEST_KNOWN_COLORS[s])

    # 3b. solved_status colors (relational column)
    if "solved_status" in headers:
        ss_idx = col_idx("solved_status")
        for ri, row in enumerate(merged, start=2):
            s = row.get("solved_status")
            if s and s in SOLVED_STATUS_COLORS:
                ws.cell(row=ri, column=ss_idx).fill = PatternFill("solid", fgColor=SOLVED_STATUS_COLORS[s])

    # 3c. hidden_bad_history colors
    if "hidden_bad_history" in headers:
        hb_idx = col_idx("hidden_bad_history")
        for ri, row in enumerate(merged, start=2):
            s = row.get("hidden_bad_history")
            if s and s in HIDDEN_BAD_COLORS:
                ws.cell(row=ri, column=hb_idx).fill = PatternFill("solid", fgColor=HIDDEN_BAD_COLORS[s])

    # 3d. cherry13b / live-arch tracking colors
    if "live_arch" in headers:
        la_idx = col_idx("live_arch")
        for ri, row in enumerate(merged, start=2):
            s = row.get("live_arch")
            if s and s in LIVE_ARCH_COLORS:
                ws.cell(row=ri, column=la_idx).fill = PatternFill("solid", fgColor=LIVE_ARCH_COLORS[s])
    if "cherry13b_tested" in headers:
        ch_idx = col_idx("cherry13b_tested")
        for ri, row in enumerate(merged, start=2):
            s = str(row.get("cherry13b_tested") or "").strip().lower()
            if s in CHERRY13B_COLORS:
                ws.cell(row=ri, column=ch_idx).fill = PatternFill("solid", fgColor=CHERRY13B_COLORS[s])
    if "unlock_potential" in headers:
        up_idx = col_idx("unlock_potential")
        for ri, row in enumerate(merged, start=2):
            s = row.get("unlock_potential")
            if s and s in UNLOCK_POTENTIAL_COLORS:
                ws.cell(row=ri, column=up_idx).fill = PatternFill("solid", fgColor=UNLOCK_POTENTIAL_COLORS[s])

    # 4. current_cost — gradient (low=green, high=red)
    cc_letter = col_letter("current_cost")
    ws.conditional_formatting.add(
        f"{cc_letter}2:{cc_letter}{last_row}",
        ColorScaleRule(
            start_type="min", start_color="63BE7B",
            mid_type="percentile", mid_value=50, mid_color="FFEB84",
            end_type="max", end_color="F8696B",
        ),
    )

    # 5. ext_savings_pct — gradient (low=green, high=red — high gap = more room to cut)
    es_letter = col_letter("ext_savings_pct")
    ws.conditional_formatting.add(
        f"{es_letter}2:{es_letter}{last_row}",
        ColorScaleRule(
            start_type="num", start_value=0, start_color="63BE7B",
            mid_type="num", mid_value=25, mid_color="FFEB84",
            end_type="num", end_value=50, end_color="F8696B",
        ),
    )

    # 6. complexity_1_10 — gradient
    cx_letter = col_letter("complexity_1_10")
    ws.conditional_formatting.add(
        f"{cx_letter}2:{cx_letter}{last_row}",
        ColorScaleRule(
            start_type="num", start_value=1, start_color="63BE7B",
            mid_type="num", mid_value=5, mid_color="FFEB84",
            end_type="num", end_value=10, end_color="F8696B",
        ),
    )

    # 7. recipe_exists — light green if True
    re_letter = col_letter("recipe_exists")
    ws.conditional_formatting.add(
        f"{re_letter}2:{re_letter}{last_row}",
        CellIsRule(operator="equal", formula=['TRUE'], fill=PatternFill("solid", fgColor="C6EFCE")),
    )

    # Re-add other sheets (Summary, Floor_Proxy, Legend) if present
    for sn, rows in other_sheets.items():
        if sn == "Tasks" or sn.startswith("Schema"):
            continue
        ws2 = wb.create_sheet(sn)
        for r in rows:
            ws2.append(list(r))

    # Add a fresh Schema sheet documenting all merged columns
    ws_schema = wb.create_sheet("Schema")
    ws_schema.append(["Column", "Source", "Description"])
    for col, src, desc in COLUMNS:
        ws_schema.append([col, src, desc])
    # Style schema header
    for ci in range(1, 4):
        cell = ws_schema.cell(row=1, column=ci)
        cell.fill = FILL_HEADER
        cell.font = FONT_HEADER
    ws_schema.column_dimensions["A"].width = 26
    ws_schema.column_dimensions["B"].width = 9
    ws_schema.column_dimensions["C"].width = 90

    wb.save(path)


def write_csv(path: Path, merged: list[dict]) -> None:
    """Regenerate task_state.csv as a flat mirror of the Tasks sheet."""
    headers = [c[0] for c in COLUMNS]
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=headers, lineterminator="\n")
        writer.writeheader()
        for row in merged:
            # Convert None -> "" for CSV cleanliness
            writer.writerow({h: ("" if row.get(h) is None else row.get(h)) for h in headers})


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser()
    p.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    p.add_argument("--csv", type=Path, default=DEFAULT_CSV)
    args = p.parse_args(argv)

    csv_rows = load_csv(args.csv)
    xlsx_rows, other_sheets = load_xlsx(args.xlsx)
    merged = merge_rows(xlsx_rows, csv_rows)
    print(f"Merging: xlsx had {len(xlsx_rows)} rows, csv had {len(csv_rows)} rows, output {len(merged)} rows")
    write_xlsx(args.xlsx, merged, other_sheets)
    write_csv(args.csv, merged)
    print(f"Wrote: {args.xlsx}")
    print(f"Wrote: {args.csv}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
