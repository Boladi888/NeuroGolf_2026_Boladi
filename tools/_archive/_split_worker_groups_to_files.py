"""Split the 10 worker groups into individual .md files under
The_400_Tasks/Task Prompts/.  Each prompt is self-contained and includes
the rebuild-from-scratch directive, multi-hour-per-task expectation,
20hr time budget, 30% adversarial rule, and the "don't give up" framing."""
import os
from pathlib import Path

import openpyxl

OUT_DIR = Path('The_400_Tasks/Task Prompts')
OUT_DIR.mkdir(parents=True, exist_ok=True)

wb = openpyxl.load_workbook('The_400_Tasks/NeuroGolf_Task_Status.xlsx')
ws = wb['Tasks']
headers = [c.value for c in ws[1]]


def ci(name):
    return headers.index(name) + 1


tasks = []
for r in range(2, ws.max_row + 1):
    tid = ws.cell(row=r, column=ci('task')).value
    fam = ws.cell(row=r, column=ci('family')).value
    transforms = ws.cell(row=r, column=ci('transformations')).value
    cost = ws.cell(row=r, column=ci('current_cost')).value or 0
    complexity = ws.cell(row=r, column=ci('complexity_1_10')).value
    if not tid:
        continue
    tasks.append({
        'tid': tid,
        'family': fam or 'Unclassified',
        'transformations': transforms or '',
        'cost': cost,
        'complexity': complexity or 0,
    })

actionable = [t for t in tasks if t['cost'] >= 1000]

GROUPS = [
    ('PR_HUGE', lambda t: t['family'] == 'Pattern_Recognition' and t['cost'] >= 50000,
     'Pattern_Recognition', 'HUGE',
     'Tile/mirror/wave repetition patterns. Tend toward Conv-heavy or per-cell-rule encodings.'),
    ('PR_BIG', lambda t: t['family'] == 'Pattern_Recognition' and 20000 <= t['cost'] < 50000,
     'Pattern_Recognition', 'BIG', 'Same family, mid-cost rewrites.'),
    ('PR_MID', lambda t: t['family'] == 'Pattern_Recognition' and 5000 <= t['cost'] < 20000,
     'Pattern_Recognition', 'MID', 'Same family, smaller wins available.'),
    ('PR_SMALL', lambda t: t['family'] == 'Pattern_Recognition' and t['cost'] < 5000,
     'Pattern_Recognition', 'SMALL',
     'Near-floor pattern tasks. DO NOT DISMISS — task299 (cost 2,678) dropped to 1,058 (60% reduction, +0.92 LB) on 2026-05-20 via bool-Concat keeping. Build prototype at The_400_Tasks/Task_by_Task/task299/prototypes/task299_bool_concat_FA59524E.onnx — study it before working any small task. Inspect every task for bool-eligible Concat/Pad chains before declaring no win.'),
    ('OB_HUGE', lambda t: t['family'] == 'Object_Based' and t['cost'] >= 50000,
     'Object_Based', 'HUGE',
     'Object segmentation / counting / coloring. Look for ReduceSum/ReduceMax bounding-box detectors and ScatterND fills.'),
    ('OB_BIG', lambda t: t['family'] == 'Object_Based' and 20000 <= t['cost'] < 50000,
     'Object_Based', 'BIG', 'Same family, mid-cost rewrites.'),
    ('OB_MID_SMALL', lambda t: t['family'] == 'Object_Based' and t['cost'] < 20000,
     'Object_Based', 'MID/SMALL', 'Smaller object-based wins.'),
    ('CL_HUGE', lambda t: t['family'] == 'Color_and_Logical' and t['cost'] >= 30000,
     'Color_and_Logical', 'HUGE',
     'Color-mapping / boolean-logic-driven tasks. Frequent Where/Equal chains; bool_andfold and or_cast_collapse often apply.'),
    ('CL_MID_SMALL', lambda t: t['family'] == 'Color_and_Logical' and t['cost'] < 30000,
     'Color_and_Logical', 'MID/SMALL', 'Smaller color/logic wins.'),
    ('SG_ALL', lambda t: t['family'] == 'Spatial_and_Geometric',
     'Spatial_and_Geometric', 'ALL',
     'Spatial transforms (rotation, reflection, tiling). Often very low cost; only big wins come from algorithm rebuilds.'),
]

FAMILY_TRICKS = {
    'Pattern_Recognition': {
        'inspiration': 'task117 (crop15+F16 -100525), task374 (conv_run_combined -11728), task381 (crop10), task233 (Equal-fold -32000), task119 (diag_conv from-scratch -80529 this session), task299 (bool-Concat 60% reduction on near-floor task)',
        'tricks': '1) Crop-to-active-grid if max_active_grid<30. 2) F16-output trick on Concat->Cast tails. 3) bool_andfold on mask Mul chains. 4) shared_cast_equal_fold on N-way color comparisons. 5) separable_mask_collapse on row/col-uniform masks.',
        'rebuild_hint': 'Many pattern tasks compute the same output via convolution OR predicate-mask OR ReduceSum-based encoding. Try the smallest-rank formulation: if the output is binary, use bool ops (1 byte/elem) throughout and Cast to F32 only at output. task119 was rebuilt from scratch this session and saved -80,529 bytes / +1.029 pts using a diagonal-conv reformulation.',
    },
    'Object_Based': {
        'inspiration': 'task266 (scalar_broadcast), task167/task103 (opset9_attr_shrink), task071 (ring-mask dynamic compute -2709 from-scratch rebuild this session)',
        'tricks': '1) ReduceMax(axes=[2]) + ReduceMax(axes=[3]) for bounding-box. 2) ScatterND for filling. 3) separable_mask_collapse on row/col masks. 4) nonzero_shortcut on Mul(input, gate)+ReduceSum chains.',
        'rebuild_hint': 'Object-based tasks often build a wide [1,C,H,W] intermediate before reducing. Try projection-cache: compute the reductions directly from input and skip the wide intermediate (see Logs/lessons.md task071 + task071-rebuild and task049 fresh-from-scratch -51523 / 96% pass). Stay at the lowest tensor rank the math allows ([H,W] not [1,1,H,W]).',
    },
    'Color_and_Logical': {
        'inspiration': 'task233 (Equal-fold -32000), task219 (Equal-fold -1056), task293 (And-fold), task196 (Or+Cast and flip_where), task368 (shared-Cast equal-fold -7345)',
        'tricks': '1) Equal-fold on Sub->Mul(self)->Less<1. 2) bool_andfold on Cast(bool)*Cast(bool) Muls. 3) or_cast_collapse on Add(Cast(gt_hi),Cast(lt_lo)). 4) flip_where on Mul(b, Sub(1, Where(c,b,0))). 5) shared_cast_equal_fold for N-way Sub-Abs-Less<1 chains.',
        'rebuild_hint': 'Stay bool through every And/Or/Not chain. Cast to F32 only at the boundary where F32 arithmetic against non-mask values is required. See Logs/lessons.md "bool-first compute" rule. The 4-byte F32 vs 1-byte bool ratio means rank-and-dtype choice can swing 4x in either direction.',
    },
    'Spatial_and_Geometric': {
        'inspiration': 'task129 (tile_to_edge_pad), task140/task087 (opset10_pad_attr), task005 (onnxsim folded), task382 (CumSum+Gather fresh rebuild this session)',
        'tricks': '1) Pad+Slice collapse via identity. 2) opset upgrade for pad-attr -> pad-input. 3) Constant-fold padding values. 4) Tile -> Reshape+Expand sometimes wins. 5) slice_chain_to_multiaxis on consecutive single-axis Slices.',
        'rebuild_hint': 'Spatial tasks often have a redundant Pad before Slice or vice versa. Look for opportunities to drop one. If the model uses Tile with a small repeat factor, try Expand instead (no element duplication). Some tasks have purely geometric semantics that can be encoded as a single CumSum+Gather (see task382).',
    },
}


def make_prompt_md(group_name, family, tier_lbl, description, members):
    tricks = FAMILY_TRICKS.get(family, FAMILY_TRICKS['Pattern_Recognition'])
    sum_cost = sum(t['cost'] for t in members)
    task_list_full = ', '.join(t['tid'] for t in members)

    # Per-task table
    table_lines = ['| task | cost | transformations | complexity |',
                   '|---|---:|---|---:|']
    for t in members:
        cmplx = t['complexity'] if t['complexity'] else '?'
        trans = (t['transformations'] or '')[:60]
        table_lines.append(f'| {t["tid"]} | {t["cost"]:,} | {trans} | {cmplx} |')
    table = '\n'.join(table_lines)

    # The full prompt — paste this into a fresh Claude Code session
    prompt = f'''```
You are a NeuroGolf 2026 Kaggle worker. Your assigned group is {group_name}.

==== GROUP CONTEXT ====
Working directory: C:\\Users\\Qx\\Documents\\Py_cx\\K-NeuroGolf
Current canonical: submissions/submission.zip @ 5835.30 LB (sha 5D9AABE0)
.venv at .venv/Scripts/python.exe

Family: {family} ({tier_lbl} cost tier)
{description}

Your tasks ({len(members)}, by cost desc):
{task_list_full}

Inspiration prototypes (look at their build_*.py for technique):
  {tricks['inspiration']}

Tricks to try first, in order:
  {tricks['tricks']}

Family-specific rebuild hint:
  {tricks['rebuild_hint']}

==== TIME + EFFORT EXPECTATIONS ====
- TOTAL BUDGET: 20 hours for this group.
- PER-TASK BUDGET: each task is a multi-hour project. Do not give up
  after 30 minutes. Try at least 3 distinct techniques per task before
  moving on. Track time spent per task in The_400_Tasks/Task_by_Task/
  taskNNN.md under "Attempts".
- DO NOT GIVE UP. Cost savings of even 1 byte translate to fractional
  LB points. The session that landed +30.11 LB today was built from
  hundreds of single-task wins.

==== WORKFLOW (per task) ====
1. Extract canonical taskNNN.onnx from submissions/submission.zip.
2. Inspect: node count, op-type histogram, value_info shapes. Look for
   30x30 intermediates (crop opportunity) and Flatten->MatMul/Gather
   (crop blocker).
3. Read inspiration prototypes' build_*.py to internalize the technique.
4. POLISH PASS: apply the tricks above in order. Score each step.
   Target: cost reduction sufficient for ~+0.1 LB gain
   (delta >= 5000 bytes on a mid-cost task is usually enough).
5. IF POLISH FAILS TO HIT +0.1 LB: this is the trigger to rebuild
   the ONNX from scratch.

   FIRST try the "input -> BOOL" technique (least invasive, biggest
   leverage): if the downstream ops are bool-tolerant (Gather,
   ReduceMax, Where, Pad, Concat, And/Or/Not, comparisons) and the
   chain doesn't include Conv/MatMul/Gemm or F32 arithmetic on the
   input, insert Cast(input -> BOOL) at the graph entry. The 4x
   dtype-width savings (F32 36k -> BOOL 9k on a [1,10,30,30] input)
   cascade through every downstream intermediate. Replace final
   Mul(gathered, mask) with And. Graph output becomes BOOL. task376
   (-18,210, +0.55 LB, verify 39/39, adversarial 100%) is the
   exemplar — read .hermes/board/staged_submits/task376_bool_t_8d70c38f_on_5835_30/manifest.json
   for the exact method.

   IF input->BOOL doesn't apply (downstream needs F32 precision),
   then try a different algorithm family:
     - Conv-heavy   -> try predicate-mask reformulation
     - Gather-heavy -> try ReduceMax-based one-hot lookup
     - ReduceSum   -> try Conv1x1 with constant weights
     - ReduceMax   -> try ArgMax + ScatterND
     - Where chain -> try direct arithmetic on integer-encoded grids
     - Static tables -> try dynamic computation at rank-2 (see task071)
   Spend several hours on the rebuild. Reference exemplars: task119
   (-80529 from-scratch), task049 (-51523 from-scratch), task071 (-2709
   from-scratch dynamic), task299 (60% reduction on near-floor),
   task376 (input->BOOL whole-graph). All landed this session.
6. AFTER THOROUGH REBUILD ATTEMPT: any negative cost delta is acceptable
   to stage, no matter how small. A 100-byte win is still a real point
   on Kaggle. Do not throw out a candidate just because it's not big.

   NEAR-FLOOR TASKS ARE NOT LOW-PRIORITY. On 2026-05-20 a 2,678-cost
   "near-floor" task dropped to 1,058 (60% reduction, +0.92 LB) by
   recognizing all Concat inputs were bool-valued and keeping them in
   bool dtype. Disproportionate wins lurk in small tasks. Always check
   for bool-eligible Concat/Pad chains before declaring no win.

==== VERIFY GATE (HARD RULE) ====
.venv/Scripts/python tools/verify_against_live.py \\
  --proto <new.onnx> --tid <NNN> --live-zip submissions/submission.zip

Required: "OK Prototype matches LIVE" AND "Prototype vs LIVE: N/N"
with no failures. This is a HARD GATE — 7 of 8 verify-failed wildcards
tested today regressed Kaggle by -13 to -17 LB. NEVER ship sub-100%
visible verify.

Single EXCEPTION: tasks with fewer than 10 visible pairs (like task000
which has 3 visible pairs) can have verify failures but still pass
hidden — verified empirically 2026-05-20. For those tasks only,
verify is advisory.

==== ADVERSARIAL GATE ====
.venv/Scripts/python tools/_score_v2.py --onnx <new.onnx> --adversarial

Threshold: >= 30% aggregate. This is the rule across the playbook
(SpreadSweepWorker.md, KanbanWorkerLoop.md). The 30-95% band is the
safe-to-ship range. Below 30% means on-contract behavior is at risk
— widen the crop window or rebuild. task162 shipped at 54.8% and
scored exactly its expected +0.22 LB, proving 30% is the right floor.

==== SCORING ====
.venv/Scripts/python tools/spread_score_isolated.py <new.onnx>

Candidate cost MUST be strictly lower than canonical taskNNN.onnx
(extract canonical via tools/spread_score_isolated.py + zipfile.open
or by inspecting submissions/submission.zip). If cost is equal or
higher, drop the candidate.

==== STAGING ====
Stage every winner under:
  .hermes/board/staged_submits/taskNNN_<technique>_t_<8hex>_on_5835_30/
    taskNNN.onnx              # the new ONNX
    manifest.json             # see SpreadSweepWorker.md step 7 template

manifest.json must include:
  - task_id, candidate_id, pattern (your technique name)
  - scores.base_cost, scores.candidate_cost, scores.cost_delta
  - tests_run including verify_against_live result and adversarial %
  - lineage_base_candidate_id pointing at the canonical 5835.30 set

==== PER-TASK NOTES ====
Append to The_400_Tasks/Task_by_Task/taskNNN.md under "Attempts":
- Date and your approximate time spent
- Techniques tried (in order)
- Verdict per technique (cost delta if applied, else why it didn't fire)
- Final staged candidate path + bytes saved (or "no candidate")
- Notes on what to try if a future worker revisits

==== COMMIT (at session end OR every 10 tasks) ====
git add -A The_400_Tasks/Task_by_Task/ .hermes/board/staged_submits/
git commit -m "fresh worker {group_name}: N staged, B bytes saved"

==== HARD RULES ====
- ONE task at a time. Independent verify + score + adversarial gates.
- DO NOT submit to Kaggle. Submit-captain has exclusive Kaggle authority.
- DO NOT edit submissions/submission.zip directly.
- DO NOT call neurogolf_sync_current_best.py.
- DO NOT skip the verify gate (except for <10-visible-pair tasks).
- DO NOT give up after one failed approach — try at least 3 per task.

==== WHEN DONE ====
Report:
  {group_name} complete: N staged candidates, B bytes saved,
  expected ~X LB combined, tasks fully exhausted: K of {len(members)},
  tasks with rebuild attempted but no candidate: M

Aim for K + M == {len(members)} (every task gets thorough effort,
even if no candidate ships).
```'''

    md = f'''# Worker prompt: `{group_name}`

**Family**: {family} ({tier_lbl} cost tier)
**Tasks**: {len(members)} | **Sum cost**: {sum_cost:,} | **Generated**: 2026-05-20 against canonical 5835.30
**Time budget**: 20 hours total | **Per-task expectation**: multi-hour project

{description}

---

## Tasks (by cost desc)

{table}

---

## Copy-paste prompt (start of fresh Claude Code session)

{prompt}
'''
    return md


for name, fltr, fam, tier_lbl, desc in GROUPS:
    members = sorted([t for t in actionable if fltr(t)], key=lambda x: -x['cost'])
    if not members:
        continue
    md = make_prompt_md(name, fam, tier_lbl, desc, members)
    out_path = OUT_DIR / f'{name}.md'
    out_path.write_text(md, encoding='utf-8')
    print(f'Wrote {out_path}  ({len(md):,} chars, {len(members)} tasks)')

print(f'\nAll 10 worker .md files written to {OUT_DIR}')
