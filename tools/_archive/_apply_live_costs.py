"""Apply live_costs.json into the xlsx (overwrite current_cost / memory / params).

Designed to be fast (single pass through the xlsx, no scoring) so it
finishes within tool timeouts. Pair with _score_zip_to_json.py.

Usage:
    python tools/_apply_live_costs.py [--xlsx XLSX] [--json JSON]
"""
import argparse, json
from pathlib import Path
import openpyxl

REPO = Path(__file__).parent.parent
DEFAULT_XLSX = REPO / 'The_400_Tasks' / 'NeuroGolf_Task_Status.xlsx'
DEFAULT_JSON = REPO / 'The_400_Tasks' / 'live_costs.json'


def main():
    p = argparse.ArgumentParser()
    p.add_argument('--xlsx', type=Path, default=DEFAULT_XLSX)
    p.add_argument('--json', type=Path, default=DEFAULT_JSON)
    args = p.parse_args()

    data = json.loads(args.json.read_text(encoding='utf-8'))
    fresh = {int(k): v for k, v in data.items()}
    print(f'Loaded {len(fresh)} fresh costs from {args.json.name}', flush=True)

    wb = openpyxl.load_workbook(args.xlsx)
    ws = wb.active
    header = [c.value for c in ws[1]]
    ti = header.index('task_id') + 1
    ci = header.index('current_cost') + 1
    mi = header.index('memory') + 1
    pi = header.index('params') + 1
    mai = header.index('macs') + 1 if 'macs' in header else None

    count = 0
    for r in range(2, ws.max_row + 1):
        tid = ws.cell(r, ti).value
        if tid is None:
            continue
        tid = int(tid)
        if tid in fresh:
            fc = fresh[tid]
            ws.cell(r, ci).value = fc['cost']
            ws.cell(r, mi).value = fc['memory']
            ws.cell(r, pi).value = fc['params']
            if mai is not None:
                ws.cell(r, mai).value = fc['macs']
            count += 1

    wb.save(args.xlsx)
    print(f'Updated {count} rows. Saved {args.xlsx.name}.', flush=True)


if __name__ == '__main__':
    main()
