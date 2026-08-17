#!/usr/bin/env python3
"""Read-only preflight for Prompter prompt-batch selection.

The helper makes the repetitive exclusion work mechanical:

- derive task ids from active prompt folders;
- check a proposed/new batch for uniqueness and prior-folder overlap;
- verify selected rows against the current workbook and team manifest;
- extract hard avoid ids from the current leaderboard note and manual flags.

It does not rank candidates or decide strategy. It only answers whether a batch
passes the cheap facts that should not cost model tokens to rediscover.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any


REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_ONNXLYFANS_ROOT = REPO_ROOT.parent / "K-NeuroGolf-OnnxlyFans"
DEFAULT_PROMPT_ROOT = REPO_ROOT / "The_400_Tasks" / "Task Prompts"
DEFAULT_WORKBOOK = DEFAULT_ONNXLYFANS_ROOT / "NG_OnnxlyFans_TaskStatus.xlsx"
DEFAULT_MANIFEST = DEFAULT_ONNXLYFANS_ROOT / "Team_Submission" / "Backend" / "manifest.csv"
DEFAULT_LEADERBOARD = (
    DEFAULT_ONNXLYFANS_ROOT / "Team_Submission" / "Backend" / "current_leaderboard_state.json"
)

TASK_RE = re.compile(r"(?i)\btask[_\s-]*0*(\d{1,3})\b|_Task0*(\d{1,3})_")
HARD_AVOID_MARKERS = (
    "hidden-failed",
    "hidden failed",
    "hidden-bad",
    "rejected",
    "regression",
    "unsafe",
    "not submitted",
    "do not submit",
    "walled",
    "hard wall",
    "proven wall",
    "known wall",
    "public-flat",
    "public flat",
)
SOFT_WALL_MARKERS = (
    "floor",
    "at floor",
    "forced floor",
    "structural floor",
    "wall",
    "no candidate",
)


def normalize_task(raw: str | int) -> str:
    text = str(raw).strip()
    match = re.search(r"(\d{1,3})", text)
    if not match:
        raise ValueError(f"cannot parse task id from {raw!r}")
    return f"task{int(match.group(1)):03d}"


def task_ids_from_text(text: str) -> set[str]:
    out: set[str] = set()
    for match in TASK_RE.finditer(text):
        number = match.group(1) or match.group(2)
        if number is not None:
            out.add(f"task{int(number):03d}")
    return out


def task_ids_from_file(path: Path) -> set[str]:
    ids = task_ids_from_text(path.name)
    if path.suffix.lower() in {".md", ".txt", ".json", ".jsonl", ".csv"}:
        try:
            ids |= task_ids_from_text(path.read_text(encoding="utf-8", errors="replace"))
        except OSError:
            pass
    return ids


def folder_sort_key(path: Path) -> tuple[int, str]:
    match = re.match(r"^(\d+)", path.name)
    return (int(match.group(1)) if match else 10_000, path.name.lower())


def resolve_prompt_folder(prompt_root: Path, token: str) -> Path:
    path = Path(token)
    return path if path.is_absolute() else prompt_root / token


def active_prompt_folders(prompt_root: Path) -> list[Path]:
    if not prompt_root.exists():
        return []
    folders = []
    for path in prompt_root.iterdir():
        if not path.is_dir() or path.name.startswith("_"):
            continue
        if re.match(r"^\d+", path.name):
            folders.append(path)
    return sorted(folders, key=folder_sort_key)


def prompt_tasks_by_folder(folders: list[Path]) -> dict[str, list[str]]:
    out: dict[str, list[str]] = {}
    for folder in folders:
        tasks: set[str] = set()
        if folder.exists():
            for path in folder.rglob("*.md"):
                # Queue filenames are the authority (`..._TaskNNN_...`). Prompt
                # bodies often mention exemplar or avoid-list task ids.
                ids = task_ids_from_text(path.name)
                tasks |= ids if ids else task_ids_from_file(path)
        out[folder.name] = sorted(tasks)
    return out


def choose_previous_folders(
    *,
    prompt_root: Path,
    new_folders: list[Path],
    explicit: list[str],
    previous_active: int,
    all_active: bool,
) -> list[Path]:
    active = active_prompt_folders(prompt_root)
    new_names = {path.name for path in new_folders}
    if explicit:
        return [resolve_prompt_folder(prompt_root, item) for item in explicit]
    eligible = [path for path in active if path.name not in new_names]
    if all_active:
        return eligible
    if new_folders:
        first_new_key = min(folder_sort_key(path) for path in new_folders)
        eligible = [path for path in eligible if folder_sort_key(path) < first_new_key]
    return eligible[-previous_active:] if previous_active > 0 else []


def read_workbook_rows(path: Path) -> tuple[dict[str, dict[str, Any]], list[str]]:
    try:
        import openpyxl  # type: ignore
    except ImportError as exc:  # pragma: no cover
        raise SystemExit("openpyxl is required to read NG_OnnxlyFans_TaskStatus.xlsx") from exc
    if not path.exists():
        raise SystemExit(f"workbook not found: {path}")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "Tasks" not in wb.sheetnames:
        raise SystemExit(f"workbook has no Tasks sheet: {path}")
    ws = wb["Tasks"]
    header = [str(cell.value).strip() if cell.value is not None else "" for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    rows: dict[str, dict[str, Any]] = {}
    for cells in ws.iter_rows(min_row=2, values_only=True):
        row = {header[i]: cells[i] for i in range(min(len(header), len(cells))) if header[i]}
        task_value = row.get("task")
        if task_value in (None, ""):
            continue
        rows[normalize_task(task_value)] = row
    return rows, header


def read_manifest_rows(path: Path) -> dict[str, dict[str, str]]:
    if not path.exists():
        raise SystemExit(f"manifest not found: {path}")
    rows: dict[str, dict[str, str]] = {}
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            if row.get("task"):
                rows[normalize_task(row["task"])] = row
    return rows


def first_value(row: dict[str, Any], names: tuple[str, ...]) -> Any:
    for name in names:
        if name in row and row[name] not in (None, ""):
            return row[name]
    return None


def extract_hard_avoid_from_text(text: str, source: str) -> dict[str, list[str]]:
    evidence: dict[str, list[str]] = {}
    for raw_line in re.split(r"[\r\n]+", text):
        line = raw_line.strip()
        if not line:
            continue
        lower = line.lower()
        if any(marker in lower for marker in HARD_AVOID_MARKERS):
            for task_id in task_ids_from_text(line):
                evidence.setdefault(task_id, []).append(f"{source}: {line[:240]}")
    return evidence


def read_leaderboard_avoid(path: Path) -> dict[str, list[str]]:
    if not path.exists():
        return {}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {}
    text_parts = []
    for key in ("note", "notes", "status", "message"):
        value = data.get(key)
        if value:
            text_parts.append(str(value))
    return extract_hard_avoid_from_text("\n".join(text_parts), str(path))


def task_doc_wall_warnings(repo_root: Path, task_ids: list[str]) -> dict[str, list[str]]:
    warnings: dict[str, list[str]] = {}
    for task_id in task_ids:
        paths = [
            repo_root / "The_400_Tasks" / "Task_by_Task" / f"{task_id}.md",
            repo_root / "The_400_Tasks" / "Task_by_Task" / task_id / "README.md",
        ]
        for path in paths:
            if not path.exists():
                continue
            try:
                text = path.read_text(encoding="utf-8", errors="replace")
            except OSError:
                continue
            for line in text.splitlines():
                lower = line.lower()
                if any(marker in lower for marker in SOFT_WALL_MARKERS):
                    warnings.setdefault(task_id, []).append(f"{path}: {line.strip()[:220]}")
                    break
    return warnings


@dataclass
class Report:
    selected_tasks: list[str]
    folders: dict[str, list[str]]
    previous_folders: dict[str, list[str]]
    hard_avoid: dict[str, list[str]]
    warnings: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    workbook_columns: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.errors


def build_report(args: argparse.Namespace) -> Report:
    prompt_root = args.prompt_root
    new_folders = [resolve_prompt_folder(prompt_root, item) for item in args.new_folders]
    selected_from_folders = prompt_tasks_by_folder(new_folders)
    selected_tasks: list[str] = []
    if args.tasks:
        raw_tasks = re.split(r"[,\s]+", " ".join(args.tasks))
        selected_tasks = sorted({normalize_task(item) for item in raw_tasks if item})
    else:
        selected_tasks = sorted({task for tasks in selected_from_folders.values() for task in tasks})

    previous_folders = choose_previous_folders(
        prompt_root=prompt_root,
        new_folders=new_folders,
        explicit=args.exclude_folders,
        previous_active=args.previous_active,
        all_active=args.all_active,
    )
    previous_by_folder = prompt_tasks_by_folder(previous_folders)
    previous_tasks = {task for tasks in previous_by_folder.values() for task in tasks}

    workbook_rows, workbook_columns = read_workbook_rows(args.workbook)
    manifest_rows = read_manifest_rows(args.manifest)

    hard_avoid = read_leaderboard_avoid(args.leaderboard)
    for item in args.avoid_task:
        task_id = normalize_task(item)
        hard_avoid.setdefault(task_id, []).append("manual --avoid-task")

    report = Report(
        selected_tasks=selected_tasks,
        folders=selected_from_folders,
        previous_folders=previous_by_folder,
        hard_avoid=hard_avoid,
        workbook_columns=workbook_columns,
    )

    if args.expect_total is not None and len(selected_tasks) != args.expect_total:
        report.errors.append(f"selected task count {len(selected_tasks)} != expected {args.expect_total}")
    if args.expect_per_folder is not None:
        for folder, tasks in selected_from_folders.items():
            if len(tasks) != args.expect_per_folder:
                report.errors.append(f"{folder} has {len(tasks)} tasks != expected {args.expect_per_folder}")
    if args.expect_counts:
        if len(args.expect_counts) != len(new_folders):
            report.errors.append(
                f"--expect-counts has {len(args.expect_counts)} entries for {len(new_folders)} --new-folders"
            )
        for folder, expected in zip(new_folders, args.expect_counts):
            tasks = selected_from_folders.get(folder.name, [])
            if len(tasks) != expected:
                report.errors.append(f"{folder.name} has {len(tasks)} tasks != expected {expected}")

    if args.tasks:
        raw_count = len([item for item in re.split(r"[,\s]+", " ".join(args.tasks)) if item])
        if raw_count != len(selected_tasks):
            report.errors.append(f"task list has duplicates: raw {raw_count}, unique {len(selected_tasks)}")

    overlap = sorted(set(selected_tasks) & previous_tasks)
    if overlap:
        report.errors.append(
            "selected tasks overlap previous prompt folders: " + ", ".join(overlap)
        )

    hard_hits = sorted(set(selected_tasks) & set(hard_avoid))
    for task_id in hard_hits:
        report.errors.append(f"{task_id} has hard avoid evidence: {' | '.join(hard_avoid[task_id])}")

    wall_warnings = task_doc_wall_warnings(REPO_ROOT, selected_tasks) if args.scan_task_docs else {}
    for task_id, evidence in sorted(wall_warnings.items()):
        message = f"{task_id} has wall/floor language in task docs: {evidence[0]}"
        if args.deny_doc_walls:
            report.errors.append(message)
        else:
            report.warnings.append(message)

    require_assigned = args.require_assigned.casefold() if args.require_assigned else ""
    for task_id in selected_tasks:
        workbook_row = workbook_rows.get(task_id)
        manifest_row = manifest_rows.get(task_id)
        if not workbook_row:
            report.errors.append(f"{task_id} missing from workbook")
            continue
        if not manifest_row:
            report.errors.append(f"{task_id} missing from manifest")
            continue
        assigned = first_value(workbook_row, ("assigned_to", "assigned", "assignee"))
        if require_assigned and str(assigned or "").casefold() != require_assigned:
            report.errors.append(f"{task_id} assigned_to={assigned!r}, expected {args.require_assigned!r}")
        workbook_cost = first_value(workbook_row, ("cost", "team_best_cost"))
        manifest_cost = manifest_row.get("cost")
        try:
            if abs(float(workbook_cost) - float(manifest_cost)) > 1e-6:
                report.errors.append(
                    f"{task_id} workbook cost {workbook_cost} != manifest cost {manifest_cost}"
                )
        except (TypeError, ValueError):
            report.errors.append(f"{task_id} has unparsable cost workbook={workbook_cost!r} manifest={manifest_cost!r}")

    return report


def render_text(report: Report) -> str:
    lines = [
        f"ok: {str(report.ok).lower()}",
        f"selected_tasks: {len(report.selected_tasks)}",
        "selected_task_ids: " + (", ".join(report.selected_tasks) if report.selected_tasks else "(none)"),
        "",
        "new_folders:",
    ]
    for folder, tasks in report.folders.items():
        lines.append(f"  - {folder}: {len(tasks)}")
    lines.append("")
    lines.append("previous_folders_used_for_overlap:")
    for folder, tasks in report.previous_folders.items():
        lines.append(f"  - {folder}: {len(tasks)}")
    if not report.previous_folders:
        lines.append("  - (none)")
    lines.append("")
    lines.append("hard_avoid_tasks:")
    for task_id, evidence in sorted(report.hard_avoid.items()):
        lines.append(f"  - {task_id}: {evidence[0]}")
    if not report.hard_avoid:
        lines.append("  - (none)")
    lines.append("")
    lines.append("errors:")
    lines.extend(f"  - {item}" for item in report.errors or ["(none)"])
    lines.append("")
    lines.append("warnings:")
    lines.extend(f"  - {item}" for item in report.warnings or ["(none)"])
    lines.append("")
    lines.append("workbook_columns: " + ", ".join(report.workbook_columns))
    return "\n".join(lines)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--prompt-root", type=Path, default=DEFAULT_PROMPT_ROOT)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    parser.add_argument("--leaderboard", type=Path, default=DEFAULT_LEADERBOARD)
    parser.add_argument("--new-folders", nargs="*", default=[], help="New prompt folder names or paths to validate")
    parser.add_argument("--tasks", nargs="*", default=[], help="Task ids to validate instead of reading --new-folders")
    parser.add_argument("--exclude-folders", nargs="*", default=[], help="Explicit previous folders for overlap checks")
    parser.add_argument("--previous-active", type=int, default=2, help="Previous active folders to exclude by default")
    parser.add_argument("--all-active", action="store_true", help="Use every active prompt folder except --new-folders")
    parser.add_argument("--avoid-task", nargs="*", default=[], help="Manual hard-exclude task ids")
    parser.add_argument("--expect-per-folder", type=int)
    parser.add_argument(
        "--expect-counts",
        nargs="*",
        type=int,
        default=[],
        help="Expected counts matching --new-folders order, e.g. --new-folders 44 45 --expect-counts 10 20",
    )
    parser.add_argument("--expect-total", type=int)
    parser.add_argument(
        "--require-assigned",
        default="",
        help="Optional assignee requirement; empty by default so prompt selection is not assignee-limited.",
    )
    parser.add_argument("--scan-task-docs", action="store_true", help="Warn on selected task docs with wall/floor language")
    parser.add_argument("--deny-doc-walls", action="store_true", help="Treat task-doc wall/floor warnings as errors")
    parser.add_argument("--json", action="store_true")
    parser.add_argument("--strict", action="store_true", help="Exit non-zero when errors are present")
    args = parser.parse_args(argv)

    report = build_report(args)
    if args.json:
        print(json.dumps(report.__dict__ | {"ok": report.ok}, indent=2, default=str))
    else:
        print(render_text(report))
    return 1 if args.strict and not report.ok else 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
