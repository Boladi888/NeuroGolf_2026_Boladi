#!/usr/bin/env python3
"""Refresh task_state.csv / NeuroGolf_Task_Status.xlsx from a promoted package.

This is the task-state half of current-best promotion. It is intentionally a
bounded, synchronous operation: score the changed task entries, patch the task
matrix, refresh derived workbook fields, and regenerate the CSV mirror. It does
not start a watcher or background process.

The normal path is automatic: tools/neurogolf_sync_current_best.py calls this
after a public-positive package promote. The CLI is for repair/backfill.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import os
import re
import sys
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile


DEFAULT_REPO_ROOT = Path(os.environ.get("NEUROGOLF_CANONICAL_ROOT", Path(__file__).resolve().parent.parent))
REPO_ROOT_FROM_FILE = Path(__file__).resolve().parent.parent
if str(REPO_ROOT_FROM_FILE) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT_FROM_FILE))
TASK_ENTRY_RE = re.compile(r"task(\d{3})\.onnx$", re.IGNORECASE)
TRAILING_BASE_RE = re.compile(r"_on_\d+(?:[._]\d+)?$", re.IGNORECASE)


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z")


def points_for_cost(cost: int | float) -> float:
    return max(1.0, 25.0 - math.log(max(1.0, float(cost))))


def source_name(candidate_id: str) -> str:
    value = (candidate_id or "").strip()
    if not value:
        return "current_best_package"
    return TRAILING_BASE_RE.sub("", value)


def task_id_from_entry(entry: str) -> int | None:
    match = TASK_ENTRY_RE.search(Path(entry).name.replace("\\", "/"))
    if not match:
        return None
    return int(match.group(1))


def normalize_task_entries(entries: str | list[str] | tuple[str, ...] | None) -> list[str]:
    if entries is None:
        values: list[str] = []
    elif isinstance(entries, str):
        values = [part.strip() for part in entries.split(",")]
    else:
        values = [str(part).strip() for part in entries]
    out = {
        f"task{tid:03d}.onnx"
        for value in values
        if (tid := task_id_from_entry(value)) is not None
    }
    return sorted(out)


def _same_path(a: Path, b: Path) -> bool:
    try:
        return a.resolve() == b.resolve()
    except OSError:
        return False


def infer_changed_task_entries(
    old_zip: Path,
    new_zip: Path,
    hinted_entries: str | list[str] | tuple[str, ...] | None = None,
) -> list[str]:
    """Return changed taskNNN.onnx entries, using byte diff plus hints.

    Hints cover single-task requested submits. The zip diff covers stacks whose
    queue metadata only says "megastack..." instead of listing every task entry.
    Any diff failure falls back to the parsed hints; promotion should not fail
    merely because the task-state convenience diff could not run.
    """
    changed = set(normalize_task_entries(hinted_entries))
    if not old_zip.exists() or not new_zip.exists() or _same_path(old_zip, new_zip):
        return sorted(changed)
    try:
        with ZipFile(old_zip) as old, ZipFile(new_zip) as new:
            old_names = set(old.namelist())
            new_names = set(new.namelist())
            for name in sorted(old_names | new_names):
                tid = task_id_from_entry(name)
                if tid is None:
                    continue
                if name not in old_names or name not in new_names:
                    changed.add(f"task{tid:03d}.onnx")
                    continue
                if old.read(name) != new.read(name):
                    changed.add(f"task{tid:03d}.onnx")
    except (BadZipFile, OSError, KeyError):
        return sorted(changed)
    return sorted(changed)


def _score_entries(zip_path: Path, entries: list[str]) -> dict[int, dict[str, Any]]:
    os.environ.setdefault("ORT_LOGGING_LEVEL", "4")
    repo_root = Path(__file__).resolve().parents[1]
    if str(repo_root / "tools") not in sys.path:
        sys.path.insert(0, str(repo_root / "tools"))
    import _score_v2 as scv2  # type: ignore

    try:
        import onnxruntime as ort  # type: ignore

        ort.set_default_logger_severity(4)
    except Exception:
        pass

    scored: dict[int, dict[str, Any]] = {}
    with ZipFile(zip_path) as zf:
        names = set(zf.namelist())
        for entry in entries:
            tid = task_id_from_entry(entry)
            if tid is None:
                continue
            if entry not in names:
                scored[tid] = {"score_status": "missing_from_zip", "entry": entry}
                continue
            tmp = tempfile.NamedTemporaryFile(suffix=".onnx", delete=False)
            try:
                tmp.write(zf.read(entry))
                tmp.close()
                result = scv2.score_network_v2(tmp.name)
            finally:
                try:
                    tmp.close()
                except OSError:
                    pass
                try:
                    os.unlink(tmp.name)
                except OSError:
                    pass
            if result is None:
                scored[tid] = {"score_status": "score_rejected", "entry": entry}
                continue
            macs, memory, params = result
            memory = int(memory or 0)
            params = int(params or 0)
            cost = memory + params
            scored[tid] = {
                "entry": entry,
                "score_status": "ok",
                "macs": int(macs or 0),
                "memory": memory,
                "params": params,
                "current_cost": cost,
                "current_points": points_for_cost(cost),
            }
    return scored


def _format_float(value: float | None, digits: int = 2) -> str:
    if value is None:
        return ""
    text = f"{value:.{digits}f}"
    return text.rstrip("0").rstrip(".")


def _relative_payload(repo_root: Path, package_path: Path, entry: str) -> str:
    try:
        rel = package_path.resolve().relative_to(repo_root.resolve()).as_posix()
    except ValueError:
        rel = package_path.as_posix()
    return f"{rel}::{entry}"


def _append_note(existing: Any, note: str) -> str:
    text = str(existing or "").strip()
    if not text:
        return note
    if note in text:
        return text
    return f"{text} | {note}"


def _load_csv_rows(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    if not path.exists():
        return [], []
    with path.open(encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        return list(reader.fieldnames or []), list(reader)


def _row_tid(row: dict[str, Any]) -> int | None:
    raw = row.get("task_id")
    if raw not in (None, ""):
        try:
            return int(raw)
        except (TypeError, ValueError):
            pass
    raw = str(row.get("task") or "")
    if raw.lower().startswith("task"):
        raw = raw[4:]
    try:
        return int(raw)
    except ValueError:
        return None


def _write_csv_rows(path: Path, fieldnames: list[str], rows: list[dict[str, Any]], *, dry_run: bool) -> None:
    if dry_run:
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, lineterminator="\n")
        writer.writeheader()
        for row in rows:
            writer.writerow({key: "" if row.get(key) is None else row.get(key) for key in fieldnames})


def _patch_csv(path: Path, updates: dict[int, dict[str, Any]], *, dry_run: bool) -> int:
    fieldnames, rows = _load_csv_rows(path)
    if not fieldnames:
        return 0
    count = 0
    for row in rows:
        tid = _row_tid(row)
        if tid not in updates:
            continue
        for key, value in updates[tid].items():
            if key in fieldnames:
                row[key] = "" if value is None else str(value)
        count += 1
    _write_csv_rows(path, fieldnames, rows, dry_run=dry_run)
    return count


def _patch_xlsx(path: Path, updates: dict[int, dict[str, Any]], *, dry_run: bool) -> int:
    import openpyxl

    wb = openpyxl.load_workbook(path)
    if "Tasks" not in wb.sheetnames:
        return 0
    ws = wb["Tasks"]
    headers = [cell.value for cell in ws[1]]
    col = {name: idx + 1 for idx, name in enumerate(headers) if name}
    if "task_id" not in col:
        return 0
    count = 0
    for row_idx in range(2, ws.max_row + 1):
        tid_value = ws.cell(row_idx, col["task_id"]).value
        try:
            tid = int(tid_value)
        except (TypeError, ValueError):
            continue
        if tid not in updates:
            continue
        for key, value in updates[tid].items():
            if key in col:
                ws.cell(row_idx, col[key]).value = value
        count += 1
    if not dry_run:
        wb.save(path)
    return count


def _latest_cost_cache(task_dir: Path) -> Path | None:
    candidates = sorted(task_dir.glob("latest_package_costs_*.csv"))
    return candidates[-1] if candidates else None


def _write_cost_cache(
    repo_root: Path,
    scored: dict[int, dict[str, Any]],
    *,
    completed_date: str,
    dry_run: bool,
) -> Path:
    task_dir = repo_root / "The_400_Tasks"
    out_path = task_dir / f"latest_package_costs_{completed_date}.csv"
    fieldnames = [
        "task",
        "task_id",
        "macs",
        "memory",
        "params",
        "current_cost",
        "current_points",
        "score_status",
    ]
    rows_by_tid: dict[int, dict[str, Any]] = {}
    latest = _latest_cost_cache(task_dir)
    if latest and latest.exists():
        with latest.open(encoding="utf-8", newline="") as handle:
            reader = csv.DictReader(handle)
            if reader.fieldnames:
                fieldnames = list(reader.fieldnames)
            for row in reader:
                tid = _row_tid(row)
                if tid is not None:
                    rows_by_tid[tid] = dict(row)
    for tid, score in scored.items():
        row = rows_by_tid.get(tid, {"task": f"task{tid:03d}", "task_id": str(tid)})
        row.update(
            {
                "task": f"task{tid:03d}",
                "task_id": str(tid),
                "macs": score.get("macs", ""),
                "memory": score.get("memory", ""),
                "params": score.get("params", ""),
                "current_cost": score.get("current_cost", ""),
                "current_points": score.get("current_points", ""),
                "score_status": score.get("score_status", ""),
            }
        )
        rows_by_tid[tid] = row
    if not dry_run:
        out_path.parent.mkdir(parents=True, exist_ok=True)
        with out_path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=fieldnames, lineterminator="\n")
            writer.writeheader()
            for tid in sorted(rows_by_tid):
                writer.writerow({key: rows_by_tid[tid].get(key, "") for key in fieldnames})
    return out_path


def _build_row_updates(
    repo_root: Path,
    package_path: Path,
    scored: dict[int, dict[str, Any]],
    *,
    candidate_id: str,
    primary_task_id: int | None = None,
    public_score: float | None,
    previous_score: float | None,
    delta: float | None,
    completed_at: str,
    mark_shipped: bool,
    existing_rows_by_tid: dict[int, dict[str, Any]] | None = None,
) -> dict[int, dict[str, Any]]:
    completed_date = completed_at[:10]
    src = source_name(candidate_id)
    score_line = ""
    if public_score is not None and previous_score is not None:
        score_line = f"{_format_float(previous_score)} -> {_format_float(public_score)}"
    elif public_score is not None:
        score_line = _format_float(public_score)
    delta_line = f"+{_format_float(delta)} LB" if delta is not None and delta >= 0 else (
        f"{_format_float(delta)} LB" if delta is not None else ""
    )
    multi = len(scored) > 1
    existing_rows_by_tid = existing_rows_by_tid or {}
    updates: dict[int, dict[str, Any]] = {}
    for tid, score in scored.items():
        entry = str(score.get("entry") or f"task{tid:03d}.onnx")
        if score.get("score_status") != "ok":
            continue
        cost = int(score["current_cost"])
        is_primary = not multi or (primary_task_id is not None and tid == primary_task_id)
        existing_row = existing_rows_by_tid.get(tid, {})
        row_src = src if is_primary else str(existing_row.get("current_src") or src)
        if is_primary:
            note = (
                f"Shipped {completed_date}: {src}, cost {cost} "
                f"(mem {score['memory']} + params {score['params']})"
            )
            if delta_line:
                note += f", {delta_line} on Kaggle"
            note += "."
        else:
            note = (
                f"Package sync {completed_date}: confirmed current package cost {cost} "
                f"(mem {score['memory']} + params {score['params']}) while promoting {src}; "
                f"preserved row source {row_src}."
            )
        row_update: dict[str, Any] = {
            "current_cost": cost,
            "current_src": row_src,
            "current_points": float(score["current_points"]),
            "macs": int(score["macs"]),
            "memory": int(score["memory"]),
            "params": int(score["params"]),
            "local_payload_path": _relative_payload(repo_root, package_path, entry),
            "reference": f"auto-sync {candidate_id or src} completed {completed_date}",
            "last_attempted_at": completed_at,
            "notes": note,
        }
        if mark_shipped:
            if is_primary:
                verdict = "Kaggle COMPLETE"
                if score_line:
                    verdict += f": {score_line}"
                if delta_line:
                    verdict += f" ({delta_line})"
                verdict += f". {src} confirmed; current shipped cost {cost}."
            else:
                verdict = f"Current package refreshed while promoting {src}; {row_src} remains shipped at cost {cost}."
            row_update.update(
                {
                    "status": "shipped",
                    "best_known_status": f"kaggle_confirmed_{completed_date}",
                    "verdict": verdict,
                    "structural_attempt": "shipped",
                    "unlock_potential": f"low_after_{cost}",
                }
            )
            if delta is not None and is_primary:
                row_update["structural_kaggle_delta"] = f"+{_format_float(delta)}"
        updates[tid] = row_update
    return updates


def _merge_notes_from_existing_csv(csv_path: Path, updates: dict[int, dict[str, Any]]) -> None:
    _, rows = _load_csv_rows(csv_path)
    existing_by_tid = {_row_tid(row): row for row in rows}
    for tid, update in updates.items():
        old = existing_by_tid.get(tid, {}).get("notes", "")
        update["notes"] = _append_note(old, str(update.get("notes") or ""))


def _run_enrich_and_merge(repo_root: Path, package_path: Path, *, dry_run: bool, run_enrich: bool) -> list[str]:
    actions: list[str] = []
    if dry_run:
        actions.append("skip enrich/merge in dry-run")
        return actions
    xlsx = repo_root / "The_400_Tasks" / "NeuroGolf_Task_Status.xlsx"
    csv_path = repo_root / "The_400_Tasks" / "task_state.csv"
    if run_enrich:
        from tools import enrich_task_state

        enrich_task_state.main(["--xlsx", str(xlsx), "--zip", str(package_path)])
        actions.append("enrich_task_state")
    from tools import merge_task_status

    merge_task_status.main(["--xlsx", str(xlsx), "--csv", str(csv_path)])
    actions.append("merge_task_status")
    return actions


def sync_task_state(
    *,
    repo_root: Path,
    package_path: Path,
    changed_entries: str | list[str] | tuple[str, ...],
    candidate_id: str = "",
    primary_task_id: int | None = None,
    public_score: float | None = None,
    previous_score: float | None = None,
    delta: float | None = None,
    completed_at: str = "",
    mark_shipped: bool = False,
    dry_run: bool = False,
    run_enrich: bool = True,
    precomputed_scores: dict[int, dict[str, Any]] | None = None,
) -> dict[str, Any]:
    repo_root = repo_root.resolve()
    package_path = package_path.resolve()
    xlsx = repo_root / "The_400_Tasks" / "NeuroGolf_Task_Status.xlsx"
    csv_path = repo_root / "The_400_Tasks" / "task_state.csv"
    if not xlsx.exists() or not csv_path.exists():
        return {
            "ok": True,
            "skipped": True,
            "reason": "task-state xlsx/csv not present in this repo root",
        }

    entries = normalize_task_entries(changed_entries)
    if not entries and precomputed_scores is None:
        return {"ok": True, "skipped": True, "reason": "no taskNNN.onnx changed entries"}
    completed_at = completed_at or utc_now()
    completed_date = completed_at[:10]

    scored = precomputed_scores if precomputed_scores is not None else _score_entries(package_path, entries)
    cache_path = _write_cost_cache(repo_root, scored, completed_date=completed_date, dry_run=dry_run)
    _, existing_rows = _load_csv_rows(csv_path)
    existing_rows_by_tid = {_row_tid(row): row for row in existing_rows}
    updates = _build_row_updates(
        repo_root,
        package_path,
        scored,
        candidate_id=candidate_id,
        primary_task_id=primary_task_id,
        public_score=public_score,
        previous_score=previous_score,
        delta=delta,
        completed_at=completed_at,
        mark_shipped=mark_shipped,
        existing_rows_by_tid=existing_rows_by_tid,
    )
    _merge_notes_from_existing_csv(csv_path, updates)

    actions: list[str] = []
    if updates:
        patched_xlsx = _patch_xlsx(xlsx, updates, dry_run=dry_run)
        patched_csv = _patch_csv(csv_path, updates, dry_run=dry_run)
        actions.extend([f"patch_xlsx_rows={patched_xlsx}", f"patch_csv_rows={patched_csv}"])
        actions.extend(_run_enrich_and_merge(repo_root, package_path, dry_run=dry_run, run_enrich=run_enrich))
    else:
        actions.append("no scoreable changed task rows")

    return {
        "ok": True,
        "skipped": False,
        "changed_entries": entries,
        "scored_task_ids": sorted(scored),
        "updated_task_ids": sorted(updates),
        "cost_cache": str(cache_path),
        "actions": actions,
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--repo-root", type=Path, default=DEFAULT_REPO_ROOT)
    parser.add_argument("--package-path", type=Path, default=None)
    parser.add_argument("--changed-entries", default="", help="Comma-separated taskNNN.onnx entries")
    parser.add_argument("--candidate-id", default="")
    parser.add_argument("--primary-task-id", type=int, default=None)
    parser.add_argument("--public-score", type=float, default=None)
    parser.add_argument("--previous-score", type=float, default=None)
    parser.add_argument("--delta", type=float, default=None)
    parser.add_argument("--completed-at", default="")
    parser.add_argument("--mark-shipped", action="store_true")
    parser.add_argument("--no-enrich", action="store_true", help="Skip derived workbook refresh; still patches costs")
    parser.add_argument("--dry-run", action="store_true")
    return parser


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    repo_root = args.repo_root.resolve()
    package_path = (args.package_path or (repo_root / "submissions" / "submission.zip")).resolve()
    result = sync_task_state(
        repo_root=repo_root,
        package_path=package_path,
        changed_entries=args.changed_entries,
        candidate_id=args.candidate_id,
        primary_task_id=args.primary_task_id,
        public_score=args.public_score,
        previous_score=args.previous_score,
        delta=args.delta,
        completed_at=args.completed_at,
        mark_shipped=args.mark_shipped,
        dry_run=args.dry_run,
        run_enrich=not args.no_enrich,
    )
    print(json.dumps(result, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
